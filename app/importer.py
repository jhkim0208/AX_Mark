"""회사 엑셀(평가대상자 · 연봉) 업로드 – 파일 읽기, 열 자동 인식, 검증, 반영.

흐름: 업로드(원본 보관) → 미리보기(열 매핑·검증 결과) → 적용(한 트랜잭션)
회사 양식이 바뀌어도 쓸 수 있도록
  - 제목 줄이 위에 있어도 머리글 행을 자동으로 찾고
  - 머리글 이름의 여러 표현(부서명/소속/조직 …)을 인식하며
  - 자동 인식이 틀리면 화면에서 열을 직접 고를 수 있게 한다.
"""
import csv
import io
import re
import secrets
from dataclasses import dataclass
from decimal import Decimal, InvalidOperation

from fastapi import HTTPException

MAX_UPLOAD_BYTES = 20 * 1024 * 1024


# ---------------------------------------------------------------------
# 열 정의 (key, 화면 이름, 필수 여부, 인식할 머리글 표현)
# ---------------------------------------------------------------------
@dataclass(frozen=True)
class Field:
    key: str
    label: str
    required: bool
    aliases: tuple[str, ...]


EMPLOYEE_FIELDS = (
    Field("dept", "부서명", True, ("부서명", "부서", "소속부서", "소속", "조직명", "조직", "팀명", "팀", "부서이름")),
    Field("emp_no", "사번", True, ("사번", "사원번호", "직원번호", "임직원번호", "개인번호", "사원코드", "empno", "id")),
    Field("name", "이름", True, ("이름", "성명", "직원명", "사원명", "임직원명", "한글성명", "name")),
    Field("level", "직급", True, ("직급", "직급명", "레벨", "cl", "직급코드", "직위", "직급레벨")),
    Field("eval_type", "평가유형", True, ("평가유형", "평가구분", "평가대상구분", "평가대상유형", "평가대상", "평가종류")),
    Field("title", "직책 (부서장 인식용)", False, ("직책", "직책명", "보직", "보직명", "직무직책")),
)

SALARY_FIELDS = (
    Field("emp_no", "사번", True, ("사번", "사원번호", "직원번호", "임직원번호", "개인번호", "사원코드", "empno", "id")),
    Field("salary", "올해 연봉", True, ("연봉", "기본연봉", "올해연봉", "당해연봉", "현재연봉", "계약연봉", "연봉총액",
                                       "연봉액", "기준연봉", "연간급여", "salary")),
    Field("incentive", "전년 인센티브 (선택)", False, ("인센티브", "성과급", "전년인센티브", "당해인센티브", "성과인센티브")),
    Field("name", "이름 (확인용)", False, ("이름", "성명", "직원명", "사원명")),
)

FIELDS = {"EMPLOYEE": EMPLOYEE_FIELDS, "SALARY": SALARY_FIELDS}
UNITS = {"원": 1, "천원": 1_000, "만원": 10_000}


def norm(text) -> str:
    return re.sub(r"[\s()\[\]{}_\-·./:]", "", str(text or "")).lower()


# ---------------------------------------------------------------------
# 파일 읽기
# ---------------------------------------------------------------------
def _cell(v):
    """엑셀 값 → 문자열/숫자 (사번 12345.0 → '12345')"""
    if v is None:
        return None
    if isinstance(v, float) and v.is_integer():
        return int(v)
    if isinstance(v, str):
        v = v.strip()
        return v or None
    return v


def read_sheets(content: bytes, file_name: str) -> dict[str, list[list]]:
    """파일의 모든 시트를 {시트명: 행 목록} 으로 읽는다."""
    lower = file_name.lower()
    if lower.endswith(".xls") and not lower.endswith(".xlsx"):
        raise HTTPException(422, "예전 엑셀 형식(.xls)은 읽을 수 없습니다. 엑셀에서 '다른 이름으로 저장 → .xlsx' 후 올려 주세요.")
    if lower.endswith(".csv"):
        for enc in ("utf-8-sig", "cp949"):
            try:
                text = content.decode(enc)
                break
            except UnicodeDecodeError:
                continue
        else:
            raise HTTPException(422, "CSV 파일의 글자 인코딩을 읽을 수 없습니다.")
        rows = [[_cell(c) for c in r] for r in csv.reader(io.StringIO(text))]
        return {"CSV": rows}
    try:
        from openpyxl import load_workbook

        wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except Exception:
        raise HTTPException(422, "엑셀 파일(.xlsx)을 열 수 없습니다. 파일이 손상되었거나 암호가 걸려 있는지 확인해 주세요.")
    sheets = {}
    for ws in wb.worksheets:
        sheets[ws.title] = [[_cell(c) for c in row] for row in ws.iter_rows(values_only=True)]
    wb.close()
    return sheets


def detect_header_row(rows: list[list], fields) -> int:
    """앞쪽 20행 중 알려진 머리글이 가장 많이 들어 있는 행 (0부터)"""
    alias_set = {norm(a) for f in fields for a in f.aliases}
    best, best_score = 0, 0
    for i, row in enumerate(rows[:20]):
        score = sum(1 for c in row if c is not None and norm(c) in alias_set)
        if score > best_score:
            best, best_score = i, score
    return best


def auto_map(headers: list, fields) -> dict[str, int | None]:
    """머리글 이름으로 열 자동 인식 (정확히 같은 이름 우선, 다음은 포함 관계)"""
    normed = [norm(h) for h in headers]
    mapping, used = {}, set()
    for f in fields:
        aliases = [norm(a) for a in f.aliases]
        idx = next((i for a in aliases for i, h in enumerate(normed) if h == a and i not in used), None)
        if idx is None:
            idx = next((i for a in aliases if len(a) >= 2 for i, h in enumerate(normed)
                        if h and a in h and i not in used), None)
        mapping[f.key] = idx
        if idx is not None:
            used.add(idx)
    return mapping


def table(content: bytes, file_name: str, kind: str, sheet: str | None, header_row: int | None) -> dict:
    sheets = read_sheets(content, file_name)
    names = list(sheets)
    if not names:
        raise HTTPException(422, "파일에 시트가 없습니다.")
    sheet = sheet if sheet in sheets else names[0]
    rows = sheets[sheet]
    if not rows:
        raise HTTPException(422, f"'{sheet}' 시트가 비어 있습니다.")
    fields = FIELDS[kind]
    hr = header_row - 1 if header_row and 1 <= header_row <= len(rows) else detect_header_row(rows, fields)
    width = max(len(r) for r in rows)
    headers = [(rows[hr][i] if i < len(rows[hr]) else None) for i in range(width)]
    headers = [str(h) if h is not None else f"(열 {i + 1})" for i, h in enumerate(headers)]
    data = []
    for offset, r in enumerate(rows[hr + 1:]):
        r = list(r) + [None] * (width - len(r))
        if all(c is None for c in r):
            continue
        data.append((hr + 2 + offset, r))                    # 엑셀 행 번호 (1부터)
    return {"sheets": names, "sheet": sheet, "header_row": hr + 1, "headers": headers, "rows": data,
            "fields": [{"key": f.key, "label": f.label, "required": f.required} for f in fields]}


# ---------------------------------------------------------------------
# 값 해석
# ---------------------------------------------------------------------
def text(v) -> str | None:
    if v is None:
        return None
    s = str(v).strip()
    return s or None


def parse_level(v) -> str | None:
    """'CL3', 'cl 3', '책임(CL3)' → 'CL3'"""
    m = re.search(r"CL\s*-?\s*([234])", str(v or "").upper())
    return f"CL{m.group(1)}" if m else None


def parse_amount(v) -> Decimal | None:
    if v is None:
        return None
    if isinstance(v, (int, float, Decimal)):
        return Decimal(str(v))
    s = re.sub(r"[,\s원₩]", "", str(v))
    try:
        return Decimal(s) if s else None
    except InvalidOperation:
        return None


def guess_unit(values: list[Decimal]) -> str:
    """연봉 숫자 크기로 단위 추정 (8,500만원 = 85,000,000 / 85,000 천원 / 8,500 만원)"""
    vals = sorted(v for v in values if v and v > 0)
    if not vals:
        return "원"
    median = vals[len(vals) // 2]
    if median >= 10_000_000:
        return "원"
    if median >= 10_000:
        return "천원"
    return "만원"


def _get(row, mapping, key):
    idx = mapping.get(key)
    return row[idx] if idx is not None and 0 <= idx < len(row) else None


def check_mapping(mapping: dict, kind: str) -> None:
    missing = [f.label for f in FIELDS[kind] if f.required and mapping.get(f.key) is None]
    if missing:
        raise HTTPException(422, f"필수 열을 선택해 주세요: {', '.join(missing)}")


# ---------------------------------------------------------------------
# 평가대상자
# ---------------------------------------------------------------------
def plan_employees(cur, rows, mapping, target_value: str, head_keyword: str | None) -> dict:
    """업로드 행 검증 + 반영 계획. DB는 바꾸지 않는다."""
    check_mapping(mapping, "EMPLOYEE")
    target_value = (target_value or "정상").strip()
    head_keyword = (head_keyword or "").strip() or None
    cur.execute("SELECT name FROM department")
    existing_depts = {r["name"] for r in cur.fetchall()}

    seen, errors, people = set(), [], []
    excluded: dict[str, int] = {}
    for line, r in rows:
        emp_no, name = text(_get(r, mapping, "emp_no")), text(_get(r, mapping, "name"))
        dept, eval_type = text(_get(r, mapping, "dept")), text(_get(r, mapping, "eval_type")) or ""
        level_raw, title = _get(r, mapping, "level"), text(_get(r, mapping, "title"))
        err = lambda reason: errors.append({"row": line, "emp_no": emp_no, "name": name, "reason": reason})
        if not emp_no:
            err("사번 없음")
            continue
        if emp_no in seen:
            err("사번 중복 (파일 안에서 두 번 이상)")
            continue
        seen.add(emp_no)
        if not name:
            err("이름 없음")
            continue
        target = eval_type == target_value
        level = parse_level(level_raw)
        if target and not dept:
            err("부서명 없음")
            continue
        if target and not level:
            err(f"직급을 CL2/CL3/CL4 로 인식할 수 없음 ('{level_raw or ''}')")
            continue
        if not target:
            excluded[eval_type or "(빈칸)"] = excluded.get(eval_type or "(빈칸)", 0) + 1
        people.append({"emp_no": emp_no, "name": name, "dept": dept, "level": level, "title": title,
                       "eval_type": eval_type or None, "target": target,
                       "is_head": bool(target and head_keyword and title and head_keyword in title)})

    depts: dict[str, dict] = {}
    for p in people:
        if not p["target"]:
            continue
        d = depts.setdefault(p["dept"], {"name": p["dept"], "new": p["dept"] not in existing_depts,
                                         "count": 0, "levels": {"CL2": 0, "CL3": 0, "CL4": 0}, "heads": []})
        d["count"] += 1
        d["levels"][p["level"]] += 1
        if p["is_head"]:
            d["heads"].append(f"{p['name']}({p['emp_no']})")

    cur.execute("SELECT count(*) AS n FROM employee WHERE eval_target AND NOT (emp_no = ANY(%s))",
                ([p["emp_no"] for p in people if p["target"]],))
    dropped = cur.fetchone()["n"]
    return {
        "total_rows": len(rows),
        "targets": sum(1 for p in people if p["target"]),
        "excluded": excluded,
        "excluded_total": sum(excluded.values()),
        "errors": errors,
        "departments": sorted(depts.values(), key=lambda d: d["name"]),
        "new_departments": sum(1 for d in depts.values() if d["new"]),
        "dropped_from_targets": dropped,     # 기존 대상자 중 이번 파일에 '정상'으로 없는 인원 → 대상 제외
        "target_value": target_value,
        "head_keyword": head_keyword,
        "people": people,
    }


def apply_employees(cur, plan: dict) -> dict:
    cur.execute("SELECT id, code FROM job_level")
    level_ids = {r["code"]: r["id"] for r in cur.fetchall()}

    # 1) 부서: 부서명으로 찾고 없으면 생성
    dept_ids = {}
    names = sorted({p["dept"] for p in plan["people"] if p["dept"]})
    for name in names:
        cur.execute("SELECT id FROM department WHERE name = %s", (name,))
        row = cur.fetchone()
        if row is None:
            cur.execute("INSERT INTO department (code, name) VALUES (%s, %s) RETURNING id",
                        (f"TMP-{secrets.token_hex(6)}", name))
            new_id = cur.fetchone()["id"]
            cur.execute("UPDATE department SET code = %s WHERE id = %s", (f"D{new_id:04d}", new_id))
            dept_ids[name] = new_id
        else:
            dept_ids[name] = row["id"]

    # 2) 직원: 사번으로 추가/갱신. 부서가 없는 대상 외 인원은 기존 직원만 대상 제외 처리
    for p in plan["people"]:
        if p["dept"]:
            cur.execute(
                """INSERT INTO employee (emp_no, name, department_id, job_level_id, job_title,
                                         eval_type, eval_target, status)
                   VALUES (%s, %s, %s, %s, %s, %s, %s, 'ACTIVE')
                   ON CONFLICT (emp_no) DO UPDATE
                      SET name = EXCLUDED.name, department_id = EXCLUDED.department_id,
                          job_level_id = COALESCE(EXCLUDED.job_level_id, employee.job_level_id),
                          job_title = EXCLUDED.job_title, eval_type = EXCLUDED.eval_type,
                          eval_target = EXCLUDED.eval_target, status = 'ACTIVE'""",
                (p["emp_no"], p["name"], dept_ids[p["dept"]], level_ids.get(p["level"]), p["title"],
                 p["eval_type"], p["target"]),
            )
        else:
            cur.execute("UPDATE employee SET eval_target = false, eval_type = %s WHERE emp_no = %s",
                        (p["eval_type"], p["emp_no"]))

    # 3) 이번 파일에 '정상'으로 없는 기존 대상자는 평가 대상에서 제외 (기록은 보존)
    targets = [p["emp_no"] for p in plan["people"] if p["target"]]
    cur.execute("UPDATE employee SET eval_target = false WHERE eval_target AND NOT (emp_no = ANY(%s))", (targets,))
    dropped = cur.rowcount

    # 4) 부서장: 직책에 키워드(예: '장')가 있는 대상자. 없으면 기존 부서장이 그대로 소속·대상이면 유지
    heads_set = 0
    for name, dept_id in dept_ids.items():
        cands = [p for p in plan["people"] if p["target"] and p["dept"] == name and p["is_head"]]
        if cands:
            cur.execute("UPDATE department SET head_employee_id = (SELECT id FROM employee WHERE emp_no = %s) "
                        "WHERE id = %s", (cands[0]["emp_no"], dept_id))
            heads_set += 1
        else:
            cur.execute(
                """UPDATE department d SET head_employee_id = NULL
                    WHERE d.id = %s AND NOT EXISTS (
                          SELECT 1 FROM employee e
                           WHERE e.id = d.head_employee_id AND e.department_id = d.id AND e.eval_target)""",
                (dept_id,),
            )
    return {"targets": len(targets), "departments": len(dept_ids), "dropped_from_targets": dropped,
            "heads_set": heads_set, "excluded_total": plan["excluded_total"], "errors": len(plan["errors"])}


# ---------------------------------------------------------------------
# 연봉
# ---------------------------------------------------------------------
def plan_salaries(cur, rows, mapping, unit: str | None, eval_year: int) -> dict:
    check_mapping(mapping, "SALARY")
    raw = [(line, text(_get(r, mapping, "emp_no")), parse_amount(_get(r, mapping, "salary")),
            parse_amount(_get(r, mapping, "incentive")) if mapping.get("incentive") is not None else None,
            text(_get(r, mapping, "name")), _get(r, mapping, "salary")) for line, r in rows]
    guessed = guess_unit([a for _, _, a, *_ in raw if a is not None])
    unit = unit if unit in UNITS else guessed
    mul = UNITS[unit]

    cur.execute("SELECT id, emp_no, name, eval_target FROM employee")
    emps = {r["emp_no"]: r for r in cur.fetchall()}
    seen, errors, not_found, matched = set(), [], [], []
    for line, emp_no, amount, inc, name, original in raw:
        if not emp_no:
            errors.append({"row": line, "emp_no": None, "name": name, "reason": "사번 없음"})
            continue
        if emp_no in seen:
            errors.append({"row": line, "emp_no": emp_no, "name": name, "reason": "사번 중복 (파일 안에서 두 번 이상)"})
            continue
        seen.add(emp_no)
        if amount is None or amount <= 0:
            errors.append({"row": line, "emp_no": emp_no, "name": name, "reason": f"연봉 금액을 읽을 수 없음 ('{original or ''}')"})
            continue
        emp = emps.get(emp_no)
        if emp is None:
            not_found.append({"row": line, "emp_no": emp_no, "name": name})
            continue
        matched.append({"employee_id": emp["id"], "emp_no": emp_no, "name": emp["name"], "target": emp["eval_target"],
                        "base_salary": int(amount * mul),
                        "incentive": int(inc * mul) if inc is not None else None,
                        "name_mismatch": bool(name and name != emp["name"])})

    have = {m["employee_id"] for m in matched}
    cur.execute(
        """SELECT e.emp_no, e.name, d.name AS dept FROM employee e JOIN department d ON d.id = e.department_id
            WHERE e.eval_target AND NOT (e.id = ANY(%s))
              AND NOT EXISTS (SELECT 1 FROM employee_compensation c
                               WHERE c.employee_id = e.id AND c.comp_year = %s)
            ORDER BY d.name, e.emp_no""",
        (list(have), eval_year),
    )
    missing = cur.fetchall()
    return {
        "unit": unit, "guessed_unit": guessed, "eval_year": eval_year,
        "total_rows": len(rows),
        "matched": len(matched),
        "matched_targets": sum(1 for m in matched if m["target"]),
        "not_found": not_found,
        "errors": errors,
        "name_mismatches": [m for m in matched if m["name_mismatch"]],
        "missing_targets": missing,                # 적용 후에도 연봉이 없는 평가 대상자
        "samples": matched[:5],
        "rows": matched,
    }


def apply_salaries(cur, plan: dict) -> dict:
    for m in plan["rows"]:
        cur.execute(
            """INSERT INTO employee_compensation (employee_id, comp_year, base_salary, incentive_amount)
               VALUES (%s, %s, %s, COALESCE(%s, 0))
               ON CONFLICT (employee_id, comp_year) DO UPDATE
                  SET base_salary = EXCLUDED.base_salary,
                      incentive_amount = COALESCE(%s, employee_compensation.incentive_amount)""",
            (m["employee_id"], plan["eval_year"], m["base_salary"], m["incentive"], m["incentive"]),
        )
    return {"updated": len(plan["rows"]), "missing_targets": len(plan["missing_targets"]),
            "not_found": len(plan["not_found"]), "errors": len(plan["errors"]), "unit": plan["unit"]}
