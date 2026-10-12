"""업로드 시연·테스트용 샘플 엑셀 생성: python scripts/make_sample_excels.py

회사 엑셀과 비슷하게 제목 줄·불필요한 열·평가유형(정상/휴직/신규입사)·오류 행을 섞어 둔다.
"""
import random
from datetime import date, timedelta
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font

OUT = Path(__file__).resolve().parents[1] / "docs" / "samples"
OUT.mkdir(parents=True, exist_ok=True)
rnd = random.Random(7)

SUR = list("김이박최정강조윤장임한오서신권황안송류홍")
GIV = ["민준", "서연", "도윤", "하은", "지호", "수아", "예준", "지우", "현우", "채원", "건우", "다은", "우진", "서윤", "준서", "유나"]
DEPTS = [("연구1팀", "팀장", 12), ("연구2팀", "팀장", 10), ("플랫폼개발실", "실장", 14)]
LEVEL_TEXT = {"CL2": ["CL2", "선임(CL2)", "cl2"], "CL3": ["CL3", "책임(CL3)", "CL 3"], "CL4": ["CL4", "수석(CL4)"]}

people, emp = [], 2019001
for dept, head_title, n in DEPTS:
    for i in range(n):
        level = "CL4" if i == 0 else rnd.choice(["CL2", "CL3", "CL3", "CL4"])
        people.append({
            "사번": emp, "성명": rnd.choice(SUR) + rnd.choice(GIV), "소속": dept,
            "직급": rnd.choice(LEVEL_TEXT[level]), "직책": head_title if i == 0 else None,
            "입사일": date(2010, 1, 1) + timedelta(days=rnd.randint(0, 5000)),
            "평가유형": "정상", "level": level,
        })
        emp += 1
# 평가 대상이 아닌 사람 / 오류 행
people[5]["평가유형"] = "휴직"
people[17]["평가유형"] = "신규입사"
people[30]["평가유형"] = "파견"
people[8]["직급"] = "수석"                                    # CL 로 인식 불가 → 오류
people.append(dict(people[3], 성명=people[3]["성명"]))         # 사번 중복 → 오류

wb = Workbook()
ws = wb.active
ws.title = "평가대상자"
ws.append(["2026년 업적평가 대상자 명단 (인사팀)"])
ws["A1"].font = Font(bold=True, size=14)
ws.append([])
cols = ["No", "사번", "성명", "소속", "직급", "직책", "입사일", "평가유형", "비고"]
ws.append(cols)
for c in ws[3]:
    c.font = Font(bold=True)
for i, p in enumerate(people, 1):
    ws.append([i, p["사번"], p["성명"], p["소속"], p["직급"], p["직책"], p["입사일"], p["평가유형"], None])
wb.create_sheet("참고").append(["평가유형: 정상 / 휴직 / 신규입사 / 파견"])
wb.save(OUT / "평가대상자_샘플.xlsx")

BASE = {"CL2": (42_000_000, 62_000_000), "CL3": (60_000_000, 88_000_000), "CL4": (85_000_000, 120_000_000)}
wb = Workbook()
ws = wb.active
ws.title = "2026 연봉"
ws.append(["사번", "성명", "부서", "직급", "2026 연봉(원)", "전년 인센티브"])
seen = set()
for p in people:
    if p["사번"] in seen:
        continue
    seen.add(p["사번"])
    lo, hi = BASE[p["level"]]
    salary = rnd.randrange(lo, hi, 100_000)
    ws.append([p["사번"], p["성명"], p["소속"], p["level"], salary, rnd.choice([0, 0, salary // 20, salary // 12])])
ws.append([2099999, "퇴사자", "연구1팀", "CL3", 70_000_000, 0])   # 시스템에 없는 사번
wb.save(OUT / "연봉_샘플.xlsx")
print("created:", *sorted(p.name for p in OUT.glob("*.xlsx")))
